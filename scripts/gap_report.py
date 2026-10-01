#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""考试配置缺口台账：①未绑定 ②已绑未发布 ③已发布但仍是「考试阅卷」
输入 = `scripts/scan_exam_cells.js` 的输出 jsonl（只读扫描产物）。

用法:
  python3 gap_report.py <cells.jsonl> <out.md> [out.xlsx]
"""
import json, os, sys, datetime
from collections import Counter, defaultdict

SUB_ORDER = ['语文', '数学', '英语', '物理', '化学', '生物', '政治', '历史', '地理']
STATUS_MEAN = {0: '未绑定试卷', 11: '已发布·考试阅卷', 15: '已发布·任务阅卷', 111: '已发布·阅卷已完成(分配锁定)'}


def main(cells_path, out_md, out_xlsx=None):
    cells = [json.loads(l) for l in open(cells_path, encoding='utf-8') if l.strip()]
    summ = {}
    sp = cells_path + '.summary.json'
    if os.path.exists(sp):
        summ = json.load(open(sp, encoding='utf-8'))

    rows = []
    for c in cells:
        if not c.get('paper'):
            kind, mode = '未绑定', '未绑定'
        elif not c.get('publishedAt'):
            kind, mode = '已绑未发布', '—'
        elif c.get('markByTask') is True:
            kind, mode = '已发布·任务阅卷', '任务阅卷'
        elif c.get('markByTask') is False:
            kind, mode = '已发布·考试阅卷', '考试阅卷'
        else:
            kind, mode = '已发布·模式未读到', '未知'
        rows.append(dict(c, kind=kind, markMode=mode,
                         pub=(c.get('publishedAt') or '')[:16].replace('T', ' ')))
    rows.sort(key=lambda x: (-int(x['exam']), SUB_ORDER.index(x['sub']) if x['sub'] in SUB_ORDER else 99))

    g1 = [r for r in rows if r['kind'] == '未绑定']
    g2 = [r for r in rows if r['kind'] == '已绑未发布']
    g3 = [r for r in rows if r['kind'] == '已发布·考试阅卷']
    ok = [r for r in rows if r['kind'] == '已发布·任务阅卷']
    unk = [r for r in rows if r['kind'] == '已发布·模式未读到']

    L = []; A = L.append
    A('# 全站考试配置缺口台账'); A('')
    A(f"- 生成时间：{datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}")
    A(f"- 范围：{len(set(r['exam'] for r in rows))} 场考试 × 9 学科 = **{len(rows)} 格**（只读扫描，未做任何写操作）")
    A(f"- 数据来源：`GET /v2/exam/<id>/subjects/` + `GET /v2/exam/subject/<id>/marker/`（直连接口，非界面点击）")
    if summ:
        A(f"- 扫描完整性：{'**完整**' if summ.get('complete') else '**不完整（有缺口，勿直接采信）**'}")
    A('')
    A('## 结论速览'); A('')
    A('| 类别 | 格数 | 说明 |'); A('|---|---|---|')
    A(f'| ① **未绑定试卷** | **{len(g1)}** | 没有试卷，无法发布、无法阅卷 |')
    A(f'| ② **已绑定但未发布** | **{len(g2)}** | 有试卷但没有发布给学生 |')
    A(f'| ③ **已发布但仍是「考试阅卷」** | **{len(g3)}** | 需要改成「任务阅卷」（不可逆） |')
    A(f'| ④ 已发布 · 任务阅卷（正常） | {len(ok)} | 含阅卷已完成的 {len([r for r in ok if r.get("status")==111])} 格 |')
    if unk:
        A(f'| ⚠️ 模式未读到 | {len(unk)} | 需人工确认 |')
    A(f'| 合计 | {len(rows)} | |'); A('')

    def dump(title, items, note=''):
        A(f'## {title}'); A('')
        if note: A(note); A('')
        if not items:
            A('**无**'); A(''); return
        A('| 考试ID | 考试名 | 学科 | 试卷ID | 发布时间 | 班级数 | 阅卷进度 |')
        A('|---|---|---|---|---|---|---|')
        for r in items:
            A(f"| {r['exam']} | {r.get('examName','')} | {r['sub']} | {r.get('paper') or ''} | {r['pub'] or '—'} | "
              f"{r.get('classroomCount') if r.get('classroomCount') is not None else '—'} | "
              f"{str(r.get('markPercent'))+'%' if r.get('markPercent') is not None else '—'} |")
        A('')
        c = Counter(r['sub'] for r in items)
        A('按学科：' + '、'.join(f'{s} {c[s]}' for s in SUB_ORDER if c.get(s))); A('')

    dump(f'① 未绑定试卷（{len(g1)} 格）', g1, '这些学科没有试卷（考试详情页显示「制作答题卡 或 关联已有试卷」），需要先补齐源试卷。')
    dump(f'② 已绑定但未发布（{len(g2)} 格）', g2, '有试卷但从未发布，学生端没有任务。')
    dump(f'③ 已发布但仍是「考试阅卷」（{len(g3)} 格）', g3, '已发布给学生，但阅卷分配是系统默认的「考试阅卷」。**改成任务阅卷后不可逆。**')
    if unk:
        dump(f'⚠️ 模式未读到（{len(unk)} 格）', unk, '需人工确认。')

    A('## 每场考试汇总'); A('')
    A('| 考试ID | 考试名 | 已绑定 | 已发布 | 未绑定 | 已绑未发布 | 仍为考试阅卷 | 任务阅卷 |')
    A('|---|---|---|---|---|---|---|---|')
    by = defaultdict(list)
    for r in rows: by[r['exam']].append(r)
    for ex in sorted(by, key=lambda e: -int(e)):
        rs = by[ex]
        A(f"| {ex} | {rs[0].get('examName','')} | {len([r for r in rs if r.get('paper')])} | "
          f"{len([r for r in rs if r.get('publishedAt')])} | {len([r for r in rs if r['kind']=='未绑定'])} | "
          f"{len([r for r in rs if r['kind']=='已绑未发布'])} | {len([r for r in rs if r['kind']=='已发布·考试阅卷'])} | "
          f"{len([r for r in rs if r['kind']=='已发布·任务阅卷'])} |")
    A('')

    os.makedirs(os.path.dirname(out_md) or '.', exist_ok=True)
    open(out_md, 'w', encoding='utf-8').write('\n'.join(L))
    print('已写', out_md)

    if out_xlsx:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        thin = Side(style='thin', color='D0D0D0'); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
        HDR = PatternFill('solid', fgColor='E8EEF7')
        FILL = {'未绑定': PatternFill('solid', fgColor='FDE2E2'),
                '考试阅卷': PatternFill('solid', fgColor='FFF4CC'),
                '任务阅卷': PatternFill('solid', fgColor='E4F5E4')}
        wb = Workbook()

        def sheet(ws, cols, data, widths, fillof=None):
            ws.append(cols)
            for i, c in enumerate(cols, 1):
                cell = ws.cell(row=1, column=i); cell.font = Font(bold=True); cell.fill = HDR
                cell.alignment = Alignment(horizontal='center', vertical='center'); cell.border = BD
            for r in data: ws.append(r)
            for ri in range(2, len(data) + 2):
                for ci in range(1, len(cols) + 1): ws.cell(row=ri, column=ci).border = BD
                if fillof:
                    f = fillof(ri - 2)
                    if f:
                        for ci in range(1, len(cols) + 1): ws.cell(row=ri, column=ci).fill = f
            for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
            ws.freeze_panes = 'A2'

        ws1 = wb.active; ws1.title = '缺口明细'
        gd = g1 + g2 + g3 + unk
        sheet(ws1, ['缺口类型', '考试ID', '考试名', '学科', '试卷ID', '发布时间', '班级数', '阅卷模式'],
              [[x['kind'], x['exam'], x.get('examName', ''), x['sub'], x.get('paper') or '（无）',
                x['pub'] or '未发布', x.get('classroomCount') if x.get('classroomCount') is not None else '', x['markMode']] for x in gd],
              [16, 9, 34, 8, 12, 18, 8, 12],
              lambda i: FILL.get('未绑定' if gd[i]['kind'] == '未绑定' else gd[i]['markMode']))

        ws2 = wb.create_sheet('全量明细')
        sheet(ws2, ['考试ID', '考试名', '学科', '试卷ID', '已发布', '发布时间', '阅卷模式', '状态码', '班级数', '学生数', '阅卷进度'],
              [[x['exam'], x.get('examName', ''), x['sub'], x.get('paper') or '（未绑定）',
                '是' if x.get('publishedAt') else '否', x['pub'], x['markMode'], x.get('status'),
                x.get('classroomCount'), x.get('studentCount'),
                (str(x.get('markPercent')) + '%') if x.get('markPercent') is not None else ''] for x in rows],
              [9, 34, 8, 12, 8, 18, 12, 8, 8, 8, 9],
              lambda i: FILL.get('未绑定' if not rows[i].get('paper') else rows[i]['markMode']))

        ws3 = wb.create_sheet('按考试汇总')
        d3 = []
        for ex in sorted(by, key=lambda e: -int(e)):
            rs = by[ex]
            d3.append([ex, rs[0].get('examName', ''), len(rs), len([r for r in rs if r.get('paper')]),
                       len([r for r in rs if r.get('publishedAt')]), len([r for r in rs if r['kind'] == '未绑定']),
                       len([r for r in rs if r['kind'] == '已绑未发布']), len([r for r in rs if r['kind'] == '已发布·考试阅卷']),
                       len([r for r in rs if r['kind'] == '已发布·任务阅卷'])])
        sheet(ws3, ['考试ID', '考试名', '学科数', '已绑定', '已发布', '未绑定', '已绑未发布', '仍为考试阅卷', '任务阅卷'],
              d3, [9, 34, 8, 8, 8, 8, 12, 13, 10])
        wb.save(out_xlsx)
        print('已写', out_xlsx)

    print(f'\n①未绑定 {len(g1)}  ②已绑未发布 {len(g2)}  ③仍考试阅卷 {len(g3)}  任务阅卷 {len(ok)}  未知 {len(unk)}')
    return 5 if unk else 0


if __name__ == '__main__':
    if len(sys.argv) < 3:
        print('用法: python3 gap_report.py <cells.jsonl> <out.md> [out.xlsx]'); sys.exit(2)
    sys.exit(main(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else None))
